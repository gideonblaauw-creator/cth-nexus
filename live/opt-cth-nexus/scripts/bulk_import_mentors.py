"""
Bulk import of mentors from CTH_Master_Mentors.xlsx into Odoo + HumHub.
Run from /opt/cth-nexus/:  python3 -m scripts.bulk_import_mentors

Skips rows with placeholder emails (pending@*, no @ sign).
Does NOT send welcome emails — temp passwords go to import_results.csv.
"""
import sys
import os
import time
import csv

sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))

from dotenv import load_dotenv
load_dotenv()

import openpyxl
from api.handlers import humhub, odoo

INPUT_FILE = "data/CTH_Master_Mentors.xlsx"
OUTPUT_FILE = "data/import_results.csv"
SLEEP_SECS = 0.5


def _is_real_email(email: str) -> bool:
    if not email or "@" not in email:
        return False
    if email.startswith("pending@"):
        return False
    return True


def main():
    wb = openpyxl.load_workbook(INPUT_FILE, data_only=True)
    ws = wb[wb.sheetnames[0]]
    headers = [c.value for c in ws[1]]
    print(f"Sheet: {wb.sheetnames[0]}  |  Headers: {headers}")
    print(f"Total rows: {ws.max_row - 1}")
    print()

    results = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        record = dict(zip(headers, row))
        email = str(record.get("Email", "") or "").strip()
        name = str(record.get("Full Name", "") or record.get("Name", "") or "").strip()

        if not _is_real_email(email):
            results.append({"id": str(record.get("ID", "")), "name": name, "email": email,
                            "status": "SKIPPED", "reason": "placeholder/missing email"})
            continue

        if not name or "TBC" in name:
            results.append({"id": str(record.get("ID", "")), "name": name, "email": email,
                            "status": "SKIPPED", "reason": "missing/TBC name"})
            continue

        data = {
            "name": name,
            "email": email.lower(),
            "phone": str(record.get("Phone / WhatsApp", "") or record.get("Phone", "") or ""),
            "linkedin": str(record.get("LinkedIn URL", "") or record.get("LinkedIn", "") or ""),
            "expertise": str(record.get("Expertise Area", "") or record.get("Expertise", "") or ""),
            "role": str(record.get("Current Role / Function", "") or record.get("Role", "") or ""),
            "company": str(record.get("Current Company", "") or record.get("Company", "") or ""),
            "programs": str(record.get("Programs", "") or ""),
            "bio": "",
            "habeas_data": "acepto",
            "source": "bulk_import",
        }

        entry = {"id": str(record.get("ID", "")), "name": name, "email": email,
                 "status": "OK", "reason": ""}

        # 1. Odoo upsert
        try:
            partner_id = odoo.upsert_mentor(data)
            entry["odoo_partner_id"] = partner_id
            print(f"  + Odoo: {name} -> partner {partner_id}")
        except Exception as e:
            entry["odoo_error"] = str(e)
            print(f"  x Odoo: {name}: {e}")

        # 2. HumHub upsert
        try:
            hh = humhub.create_or_update_user(data)
            entry["humhub_user_id"] = hh["user_id"]
            entry["humhub_username"] = hh["username"]
            entry["temp_password"] = hh["temp_password"]
            print(f"  + HumHub: {name} -> {hh['username']}")
        except Exception as e:
            entry["humhub_error"] = str(e)
            entry["status"] = "PARTIAL" if "odoo_partner_id" in entry else "ERROR"
            print(f"  x HumHub: {name}: {e}")

        if "odoo_error" in entry and "humhub_error" in entry:
            entry["status"] = "ERROR"
        elif "odoo_error" in entry or "humhub_error" in entry:
            entry["status"] = "PARTIAL"

        results.append(entry)
        time.sleep(SLEEP_SECS)

    # Write results CSV
    if results:
        all_keys = set()
        for r in results:
            all_keys.update(r.keys())
        fieldnames = ["id", "name", "email", "status", "reason",
                      "odoo_partner_id", "odoo_error",
                      "humhub_user_id", "humhub_username", "temp_password", "humhub_error"]
        fieldnames = [k for k in fieldnames if k in all_keys]
        extra = sorted(all_keys - set(fieldnames))
        fieldnames.extend(extra)

        os.makedirs(os.path.dirname(OUTPUT_FILE) or ".", exist_ok=True)
        with open(OUTPUT_FILE, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
            writer.writeheader()
            writer.writerows(results)

    ok = sum(1 for r in results if r.get("status") == "OK")
    partial = sum(1 for r in results if r.get("status") == "PARTIAL")
    skipped = sum(1 for r in results if r.get("status") == "SKIPPED")
    error = sum(1 for r in results if r.get("status") == "ERROR")
    print(f"\nDone: {ok} OK, {partial} partial, {skipped} skipped, {error} errors")
    print(f"Results saved to {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
